"""Hitung Cohen's kappa antar dua pelabel + susun daftar ketidaksepakatan untuk adjudikasi.

Dipakai pada Tahap 2 rubrik (setelah pilot 50 teks) dan diulang setiap batch pelabelan berikutnya.
Ambang lanjut: kappa >= 0,61 (Landis & Koch: "kuat").

Perlakuan `tidak_dapat_dinilai`: baris ini **dikeluarkan** dari perhitungan kappa, karena ia bukan
kelas model melainkan penanda proses — memasukkannya akan menaikkan kappa secara semu lewat
kesepakatan pada baris yang justru tidak akan dipakai melatih. Kesepakatan atas keputusan
"dapat/tidak dapat dinilai" dilaporkan terpisah.

Pemakaian:
    python progres/fase-1-data-pelabelan/hitung_kappa.py \
        --a lembar-kerja-pelabelan-A.xlsx --b lembar-kerja-pelabelan-B.xlsx
"""
from __future__ import annotations

import argparse
import csv
from collections import Counter
from pathlib import Path

AMBANG_LANJUT = 0.61
TDN = "tidak_dapat_dinilai"
KELAS = ["rendah", "tinggi"]

TAFSIR = [
    (0.00, "tidak ada kesepakatan", "Definisi bermasalah — susun ulang bersama pembimbing."),
    (0.20, "sangat lemah", "Hentikan pelabelan; tinjau ulang rubrik menyeluruh."),
    (0.40, "lemah", "Hentikan; perjelas definisi TINGGI dan tabel kasus batas."),
    (0.60, "sedang", "Belum cukup. Revisi rubrik, ulangi pilot dengan 50 teks BARU."),
    (0.80, "kuat", "Boleh lanjut ke pelabelan penuh."),
    (1.01, "sangat kuat", "Boleh lanjut ke pelabelan penuh."),
]


def tafsirkan(k: float) -> tuple[str, str]:
    for batas, label, tindakan in TAFSIR:
        if k <= batas:
            return label, tindakan
    return TAFSIR[-1][1], TAFSIR[-1][2]


def muat(path: Path) -> dict[str, dict[str, str]]:
    """Baca lembar kerja (.xlsx sheet 'Pelabelan', atau .csv) → {id_teks: {teks, label, catatan}}."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=True)
        ws = wb["Pelabelan"] if "Pelabelan" in wb.sheetnames else wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        header = [str(h).strip() if h else "" for h in rows[0]]
        isi = [dict(zip(header, r)) for r in rows[1:]]
    else:
        with path.open(encoding="utf-8", newline="") as f:
            isi = list(csv.DictReader(f))

    hasil: dict[str, dict[str, str]] = {}
    for r in isi:
        kode = str(r.get("id_teks") or "").strip()
        if not kode:
            continue
        hasil[kode] = {
            "teks": str(r.get("teks_anonim") or r.get("teks") or "").strip(),
            "label": str(r.get("label") or "").strip().lower(),
            "catatan": str(r.get("catatan") or "").strip(),
        }
    if not hasil:
        raise SystemExit(f"Tidak ada baris terbaca dari {path}")
    return hasil


def tulis_ketidaksepakatan(baris: list[dict[str, str]], out: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Adjudikasi"
    kolom = ["id_teks", "teks_anonim", "label_A", "label_B", "catatan_A", "catatan_B",
             "label_final", "alasan_putusan"]
    lebar = [12, 76, 14, 14, 30, 30, 16, 40]
    for i, (judul, w) in enumerate(zip(kolom, lebar), start=1):
        sel = ws.cell(row=1, column=i, value=judul)
        sel.font = Font(bold=True, color="1E3A5F")
        sel.fill = PatternFill("solid", fgColor="EBEFF5")
        ws.column_dimensions[get_column_letter(i)].width = w

    dv = DataValidation(type="list", formula1=f'"{",".join(KELAS + [TDN])}"', allow_blank=True)
    ws.add_data_validation(dv)

    for i, b in enumerate(baris, start=2):
        for j, k in enumerate(kolom, start=1):
            sel = ws.cell(row=i, column=j, value=b.get(k, ""))
            sel.alignment = Alignment(wrap_text=(k == "teks_anonim"), vertical="top")
        final = ws.cell(row=i, column=7)
        final.fill = PatternFill("solid", fgColor="FFF9DB")
        dv.add(final)
        ws.row_dimensions[i].height = 30
    ws.freeze_panes = "A2"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description="Cohen's kappa antar dua pelabel + daftar adjudikasi.")
    ap.add_argument("--a", required=True, help="lembar kerja pelabel A (.xlsx atau .csv)")
    ap.add_argument("--b", required=True, help="lembar kerja pelabel B")
    ap.add_argument("--out", default=None, help="berkas keluaran daftar ketidaksepakatan (.xlsx)")
    args = ap.parse_args()

    from sklearn.metrics import cohen_kappa_score, confusion_matrix

    A, B = muat(Path(args.a)), muat(Path(args.b))

    hanya_a, hanya_b = set(A) - set(B), set(B) - set(A)
    if hanya_a or hanya_b:
        print(f"PERINGATAN: id_teks tidak sama persis — hanya di A: {len(hanya_a)}, "
              f"hanya di B: {len(hanya_b)}. Hanya irisannya yang dihitung.\n")
    kode = sorted(set(A) & set(B))

    kosong = [k for k in kode if not A[k]["label"] or not B[k]["label"]]
    if kosong:
        print(f"PERINGATAN: {len(kosong)} baris belum berlabel di salah satu lembar, dilewati.\n")
    terisi = [k for k in kode if k not in kosong]

    dipakai = {A[k]["label"] for k in terisi} | {B[k]["label"] for k in terisi}
    tidak_sah = dipakai - set(KELAS) - {TDN}
    if tidak_sah:
        raise SystemExit(f"Label tidak dikenal: {sorted(tidak_sah)}. Sah: {KELAS + [TDN]}")

    # --- kesepakatan atas keputusan "dapat dinilai" ---
    tdn_a = {k for k in terisi if A[k]["label"] == TDN}
    tdn_b = {k for k in terisi if B[k]["label"] == TDN}
    dinilai = [k for k in terisi if k not in tdn_a and k not in tdn_b]

    print(f"Berkas A : {args.a}")
    print(f"Berkas B : {args.b}")
    print(f"\nBaris beririsan & berlabel : {len(terisi)}")
    print(f"  'tidak_dapat_dinilai' A  : {len(tdn_a)}")
    print(f"  'tidak_dapat_dinilai' B  : {len(tdn_b)}")
    print(f"  sepakat 'tidak dapat dinilai': {len(tdn_a & tdn_b)}")
    print(f"  dipakai menghitung kappa : {len(dinilai)}")

    if len(dinilai) < 20:
        print("\nPERINGATAN: kurang dari 20 baris — kappa tidak stabil pada sampel sekecil ini.")
    if not dinilai:
        raise SystemExit("Tidak ada baris yang dapat dihitung.")

    ya, yb = [A[k]["label"] for k in dinilai], [B[k]["label"] for k in dinilai]
    sepakat = sum(1 for x, y in zip(ya, yb) if x == y)
    kappa = float(cohen_kappa_score(ya, yb, labels=KELAS))
    label_tafsir, tindakan = tafsirkan(kappa)

    print(f"\nKesepakatan mentah : {sepakat}/{len(dinilai)} ({sepakat / len(dinilai):.1%})")
    print(f"Cohen's kappa      : {kappa:.4f}  ({label_tafsir})")
    print(f"Ambang lanjut      : {AMBANG_LANJUT}")
    print(f"Putusan            : {'LANJUT' if kappa >= AMBANG_LANJUT else 'BELUM BOLEH LANJUT'}")
    print(f"Tindakan           : {tindakan}")

    print("\nMatriks kesepakatan (baris = A, kolom = B):")
    m = confusion_matrix(ya, yb, labels=KELAS)
    print(f"{'':>10}" + "".join(f"{k:>10}" for k in KELAS))
    for i, k in enumerate(KELAS):
        print(f"{k:>10}" + "".join(f"{m[i][j]:>10}" for j in range(len(KELAS))))

    print("\nSebaran label:")
    for nama, y in (("A", ya), ("B", yb)):
        c = Counter(y)
        print(f"  {nama}: " + ", ".join(f"{k}={c.get(k, 0)}" for k in KELAS))

    # --- daftar ketidaksepakatan (termasuk selisih soal 'tidak dapat dinilai') ---
    beda = [k for k in terisi if A[k]["label"] != B[k]["label"]]
    print(f"\nKetidaksepakatan: {len(beda)} baris")
    if beda:
        out = Path(args.out) if args.out else Path(args.a).parent / "ketidaksepakatan.xlsx"
        baris = [
            {"id_teks": k, "teks_anonim": A[k]["teks"] or B[k]["teks"],
             "label_A": A[k]["label"], "label_B": B[k]["label"],
             "catatan_A": A[k]["catatan"], "catatan_B": B[k]["catatan"],
             "label_final": "", "alasan_putusan": ""}
            for k in beda
        ]
        print(f"Daftar adjudikasi ditulis ke: {tulis_ketidaksepakatan(baris, out)}")
        print("\nLangkah berikut (Tahap 4 rubrik): diskusikan berdua, isi kolom label_final.")
        print("Bila tetap tidak sepakat, putusan diserahkan kepada pihak ketiga.")
        print("Putusan atas kasus yang belum tercakup dicatat sebagai amandemen rubrik Bagian 10.")
    else:
        print("Kedua pelabel sepakat pada seluruh baris.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
